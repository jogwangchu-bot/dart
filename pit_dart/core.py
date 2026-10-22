from __future__ import annotations

import csv
import hashlib
import io
import json
import os
import re
import tempfile
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

from openpyxl import load_workbook


class CollectorError(Exception):
    """Messages are controlled locally, never raw request exceptions or URLs."""
    code = "FETCH_FAILED"


class InputError(CollectorError):
    code = "INPUT_INVALID"


class ReviewRequired(CollectorError):
    code = "REVIEW_REQUIRED"


class IntegrityError(CollectorError):
    code = "INTEGRITY_FAILED"


class DailyLimit(CollectorError):
    code = "DAILY_LIMIT"


class CredentialError(CollectorError):
    code = "CREDENTIAL_REQUIRED"


class APIHTTPError(CollectorError):
    """Structured HTTP diagnostics, without request URLs or credentials."""
    def __init__(self, endpoint: str, http_status: int, exhausted: bool = False):
        if endpoint not in ("corpCode.xml", "list.json", "document.xml") or not 100 <= http_status <= 599:
            raise ValueError("Invalid API HTTP diagnostic")
        self.endpoint, self.http_status = endpoint, http_status
        suffix = " after bounded retries" if exhausted else ""
        super().__init__(f"OpenDART {endpoint}: HTTP {http_status}{suffix}")


def digest(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def canonical(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"))


def atomic_write(path: Path, data: bytes) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    descriptor, staging = tempfile.mkstemp(prefix=f".{path.name}.", dir=path.parent)
    try:
        with os.fdopen(descriptor, "wb") as stream:
            stream.write(data)
            stream.flush()
            os.fsync(stream.fileno())
        os.replace(staging, path)
    finally:
        if os.path.exists(staging):
            os.unlink(staging)


def write_json(path: Path, value: Any) -> None:
    atomic_write(path, (canonical(value) + "\n").encode("utf-8"))


def write_csv(path: Path, fields: list[str], rows: list[dict]) -> None:
    stream = io.StringIO(newline="")
    writer = csv.DictWriter(stream, fieldnames=fields, extrasaction="ignore", lineterminator="\n")
    writer.writeheader()
    writer.writerows({key: row.get(key, "") for key in fields} for row in rows)
    atomic_write(path, stream.getvalue().encode("utf-8-sig"))


def cutoff(asof: date) -> datetime:
    return datetime.combine(asof, time.min, ZoneInfo("Asia/Seoul")) - timedelta(seconds=1)


def receipt_url(number: str) -> str:
    if not re.fullmatch(r"\d{14}", number):
        raise IntegrityError("Invalid receipt number")
    return f"https://dart.fss.or.kr/dsaf001/main.do?rcpNo={number}"


@dataclass(frozen=True)
class Stock:
    order: int
    code: str
    row: int
    # These are administrative hints, NEVER evaluation fields.
    hints: tuple[str, ...] = field(default_factory=tuple, repr=False)


def read_universe(path: Path, sheet: str | None = None, expected_count: int | None = None) -> list[Stock]:
    try:
        workbook = load_workbook(path, read_only=True, data_only=False)
    except Exception:
        raise InputError("Cannot open input workbook") from None
    result, seen = [], set()
    try:
        selected = workbook[sheet] if sheet else workbook.worksheets[0]
        header = next(selected.iter_rows(values_only=True), ())
        if len(header) < 3 or str(header[0]).strip() != "순번" or str(header[2]).strip() != "종목코드":
            raise InputError("Expected A=순번 and C=종목코드; select the correct sheet")
        previous_order = 0
        for row_number, row in enumerate(selected.iter_rows(min_row=2, values_only=True), 2):
            if all(cell is None for cell in row):
                continue
            if len(row) < 3 or isinstance(row[0], bool) or not isinstance(row[0], int):
                raise InputError(f"Invalid sequence at worksheet row {row_number}")
            order = row[0]
            if order <= previous_order:
                raise InputError(f"Sequence is not strictly increasing at row {row_number}; input not reordered")
            value = row[2]
            if isinstance(value, int) and not isinstance(value, bool) and 0 <= value <= 999999:
                code = f"{value:06d}"
            elif isinstance(value, str) and re.fullmatch(r"\d{6}", value.strip(), flags=re.ASCII):
                code = value.strip()
            else:
                raise InputError(f"Invalid six-digit stock code at row {row_number}")
            if code in seen:
                raise InputError(f"Duplicate stock code at row {row_number}")
            seen.add(code)
            result.append(Stock(order, code, row_number, tuple(str(row[i] or "") for i in (1, 3, 4) if i < len(row))))
            previous_order = order
    except KeyError:
        raise InputError("Requested worksheet not found") from None
    finally:
        workbook.close()
    if not result or (expected_count is not None and len(result) != expected_count):
        raise InputError(f"Unexpected universe size: {len(result)}")
    return result


def universe_projection(stocks: list[Stock], asof: date, verified_names: dict | None = None) -> list[dict]:
    names = verified_names or {}
    return [{"order": stock.order, "stock_code": stock.code, "asof": asof.isoformat(),
             "historical_name": names.get(stock.code, "")} for stock in stocks]
